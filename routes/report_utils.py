"""
report_utils.py — utilitários compartilhados para relatórios e filtros.

Fornece:
  • resolução de escopo de EDOTs respeitando o perfil do usuário e os
    filtros opcionais de OPO / Hospital;
  • parsing de período (data_inicio / data_fim, com fallback para `dias`);
  • geradores de arquivo CSV, XLSX e PDF a partir de linhas genéricas.
"""
import csv
import io
from datetime import datetime, timedelta

from flask import Response
from sqlalchemy import func, distinct

from models import db, EDOT, OPO, Paciente, PacienteHistorico


# ─────────────────────────────────────────────────────────────────────────────
# Contagem de funil clínico (coorte por data de cadastro do paciente)
# ─────────────────────────────────────────────────────────────────────────────
# Um paciente é "notificado de M.E." se em algum momento esteve em qualquer
# status da fase de M.E. — seja no cadastro (campo 'criacao') ou numa transição
# (campo 'status'). Os desfechos terminais são subconjuntos dessa fase, então o
# funil sempre fecha: notificações ≥ (efetivação + PCR + NAF + CIM + em andamento).
ME_STAGE_STATUSES = (
    'protocolo_me',
    'me_sem_confirmacao',
    'me_confirmado',
    'me_com_doacao',
    'me_cim',
    'me_naf',
    'pcr_antes_doacao',
)


def contar_pacientes_status(edot_ids, statuses, desde=None, ate=None):
    """Conta pacientes distintos que já atingiram QUALQUER um dos `statuses`.

    Considera tanto o status de cadastro ('criacao') quanto transições
    ('status'), evitando subcontagem de pacientes já cadastrados em protocolo.
    O período filtra a COORTE pela data de cadastro do paciente
    (Paciente.created_at), garantindo que o funil feche dentro do intervalo.
    """
    if isinstance(statuses, str):
        statuses = (statuses,)
    q = (
        db.session.query(func.count(distinct(Paciente.id)))
        .join(PacienteHistorico, PacienteHistorico.paciente_id == Paciente.id)
        .filter(
            Paciente.edot_id.in_(edot_ids),
            PacienteHistorico.campo_alterado.in_(('status', 'criacao')),
            PacienteHistorico.valor_novo.in_(tuple(statuses)),
        )
    )
    if desde:
        q = q.filter(Paciente.created_at >= desde)
    if ate:
        q = q.filter(Paciente.created_at <= ate)
    return q.scalar() or 0



# ─────────────────────────────────────────────────────────────────────────────
# Escopo / filtros
# ─────────────────────────────────────────────────────────────────────────────
def resolver_edot_ids(claims, opo_id=None, edot_id=None):
    """Retorna a lista de edot_ids que o usuário pode ver, aplicando os
    filtros opcionais de OPO e Hospital sem jamais ampliar o escopo do perfil.

    - cet_admin: todas as EDOTs; pode filtrar por opo_id e/ou edot_id
    - opo:       apenas EDOTs da sua OPO; pode filtrar por edot_id (dentro dela)
    - edot_*:    apenas a própria EDOT (filtros ignorados)
    """
    perfil = claims.get('perfil')

    if perfil == 'cet_admin':
        q = EDOT.query
        if opo_id:
            q = q.filter_by(opo_id=opo_id)
        if edot_id:
            q = q.filter_by(id=edot_id)
        return [e.id for e in q.all()]

    if perfil == 'opo':
        q = EDOT.query.filter_by(opo_id=claims.get('opo_id'))
        if edot_id:
            q = q.filter_by(id=edot_id)
        return [e.id for e in q.all()]

    # edot_coord / edot_membro
    proprio = claims.get('edot_id')
    return [proprio] if proprio else []


def _parse_data(valor, fim_do_dia=False):
    """Converte 'YYYY-MM-DD' (ou ISO) em datetime. Retorna None se inválido."""
    if not valor:
        return None
    for fmt in ('%Y-%m-%d', '%Y-%m-%dT%H:%M', '%Y-%m-%dT%H:%M:%S'):
        try:
            dt = datetime.strptime(valor, fmt)
            if fim_do_dia and fmt == '%Y-%m-%d':
                dt = dt.replace(hour=23, minute=59, second=59)
            return dt
        except ValueError:
            continue
    return None


def resolver_periodo(args):
    """Determina o intervalo [desde, ate] a partir dos args da requisição.

    Prioridade: data_inicio/data_fim explícitos > `dias` > None.
    Retorna (desde, ate, rotulo) onde desde/ate podem ser None (= sem limite).
    """
    data_inicio = _parse_data(args.get('data_inicio'))
    data_fim    = _parse_data(args.get('data_fim'), fim_do_dia=True)

    if data_inicio or data_fim:
        rotulo = 'Período: {} a {}'.format(
            data_inicio.strftime('%d/%m/%Y') if data_inicio else 'início',
            data_fim.strftime('%d/%m/%Y') if data_fim else 'hoje',
        )
        return data_inicio, data_fim, rotulo

    dias = args.get('dias', type=int)
    if dias:
        desde = datetime.utcnow() - timedelta(days=dias)
        return desde, None, f'Últimos {dias} dias'

    return None, None, 'Período completo'


def rotulo_escopo(claims, opo_id=None, edot_id=None):
    """Texto legível descrevendo o escopo aplicado (para subtítulo do relatório)."""
    partes = []
    if opo_id:
        opo = OPO.query.get(opo_id)
        if opo:
            partes.append(f'OPO: {opo.nome}')
    if edot_id:
        edot = EDOT.query.get(edot_id)
        if edot:
            partes.append(f'Hospital: {edot.hospital_nome}')
    if not partes:
        perfil = claims.get('perfil')
        if perfil == 'cet_admin':
            partes.append('Todas as OPOs e hospitais')
        elif perfil == 'opo':
            partes.append('Todos os hospitais da OPO')
    return ' · '.join(partes)


# ─────────────────────────────────────────────────────────────────────────────
# Geradores de arquivo
# ─────────────────────────────────────────────────────────────────────────────
def _valor(linha, chave):
    v = linha.get(chave)
    return '' if v is None else str(v)


def gerar_csv(nome_arquivo, colunas, linhas):
    """colunas: lista de (chave, cabeçalho). linhas: lista de dicts."""
    buffer = io.StringIO()
    buffer.write('﻿')  # BOM p/ acentuação correta no Excel
    writer = csv.writer(buffer, delimiter=';')
    writer.writerow([cab for _, cab in colunas])
    for linha in linhas:
        writer.writerow([_valor(linha, ch) for ch, _ in colunas])

    return Response(
        buffer.getvalue(),
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': f'attachment; filename="{nome_arquivo}.csv"'},
    )


def gerar_xlsx(nome_arquivo, titulo, subtitulo, colunas, linhas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Relatório'

    azul = PatternFill(start_color='1B5E9B', end_color='1B5E9B', fill_type='solid')
    branco_bold = Font(bold=True, color='FFFFFF')

    linha_atual = 1
    # Título
    ws.cell(row=linha_atual, column=1, value=titulo).font = Font(bold=True, size=14)
    linha_atual += 1
    if subtitulo:
        ws.cell(row=linha_atual, column=1, value=subtitulo).font = Font(
            italic=True, size=10, color='666666'
        )
        linha_atual += 1
    ws.cell(row=linha_atual, column=1,
            value='Emitido em ' + datetime.utcnow().strftime('%d/%m/%Y')).font = Font(
        size=9, color='999999'
    )
    linha_atual += 2

    # Cabeçalho
    cab_row = linha_atual
    for c, (_, cab) in enumerate(colunas, start=1):
        cel = ws.cell(row=cab_row, column=c, value=cab)
        cel.fill = azul
        cel.font = branco_bold
        cel.alignment = Alignment(horizontal='center')
    linha_atual += 1

    # Dados
    for linha in linhas:
        for c, (ch, _) in enumerate(colunas, start=1):
            ws.cell(row=linha_atual, column=c, value=linha.get(ch))
        linha_atual += 1

    # Largura das colunas
    for c, (ch, cab) in enumerate(colunas, start=1):
        largura = max(len(cab), *(len(_valor(l, ch)) for l in linhas)) if linhas else len(cab)
        ws.column_dimensions[get_column_letter(c)].width = min(max(largura + 2, 10), 45)

    ws.freeze_panes = ws.cell(row=cab_row + 1, column=1)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return Response(
        out.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{nome_arquivo}.xlsx"'},
    )


def gerar_pdf(nome_arquivo, titulo, subtitulo, colunas, linhas):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
    )

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=landscape(A4),
        leftMargin=1 * cm, rightMargin=1 * cm,
        topMargin=1 * cm, bottomMargin=1 * cm,
    )
    estilos = getSampleStyleSheet()
    est_titulo = ParagraphStyle('t', parent=estilos['Title'], fontSize=16,
                                textColor=colors.HexColor('#1B5E9B'), spaceAfter=4)
    est_sub = ParagraphStyle('s', parent=estilos['Normal'], fontSize=9,
                             textColor=colors.HexColor('#666666'))
    est_cel = ParagraphStyle('c', parent=estilos['Normal'], fontSize=7, leading=9)
    est_cab = ParagraphStyle('h', parent=estilos['Normal'], fontSize=7.5,
                             leading=9, textColor=colors.white, fontName='Helvetica-Bold')

    elementos = [Paragraph(titulo, est_titulo)]
    if subtitulo:
        elementos.append(Paragraph(subtitulo, est_sub))
    elementos.append(Paragraph(
        'Emitido em ' + datetime.utcnow().strftime('%d/%m/%Y'), est_sub))
    elementos.append(Spacer(1, 0.4 * cm))

    dados = [[Paragraph(cab, est_cab) for _, cab in colunas]]
    for linha in linhas:
        dados.append([Paragraph(_valor(linha, ch).replace('\n', '<br/>'), est_cel)
                      for ch, _ in colunas])

    tabela = Table(dados, repeatRows=1)
    tabela.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B5E9B')),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1),
         [colors.white, colors.HexColor('#F2F6FA')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elementos.append(tabela)

    if not linhas:
        elementos.append(Spacer(1, 0.5 * cm))
        elementos.append(Paragraph('Nenhum registro encontrado para os filtros selecionados.', est_sub))

    doc.build(elementos)
    out.seek(0)
    return Response(
        out.getvalue(),
        mimetype='application/pdf',
        headers={'Content-Disposition': f'attachment; filename="{nome_arquivo}.pdf"'},
    )


def gerar_relatorio(formato, nome_arquivo, titulo, subtitulo, colunas, linhas):
    """Dispatcher: formato ∈ {'csv','xlsx','pdf'}."""
    if formato == 'csv':
        return gerar_csv(nome_arquivo, colunas, linhas)
    if formato == 'xlsx':
        return gerar_xlsx(nome_arquivo, titulo, subtitulo, colunas, linhas)
    if formato == 'pdf':
        return gerar_pdf(nome_arquivo, titulo, subtitulo, colunas, linhas)
    return None
